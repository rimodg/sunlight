#!/usr/bin/env python3
"""Generate SUNLIGHT Package README (French) as .docx."""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, 'deliverables', 'SUNLIGHT_Package_README_FR.docx')

NAVY = RGBColor(0x1A, 0x2A, 0x44)
DARK_GREY = RGBColor(0x33, 0x33, 0x33)

doc = Document()

# Page setup
for section in doc.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.25)
    section.bottom_margin = Inches(1.25)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Georgia'
font.size = Pt(11)
font.color.rgb = DARK_GREY

for level in range(1, 4):
    h_style = doc.styles[f'Heading {level}']
    h_font = h_style.font
    h_font.name = 'Georgia'
    h_font.color.rgb = NAVY
    h_font.bold = True
    if level == 1:
        h_font.size = Pt(16)
    elif level == 2:
        h_font.size = Pt(13)
    else:
        h_font.size = Pt(11)


def add_accent_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="4" w:space="1" w:color="1A2A44"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


# ── Title ──
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(48)
run = p.add_run('SUNLIGHT')
run.font.name = 'Georgia'
run.font.size = Pt(22)
run.font.color.rgb = NAVY
run.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
run = p.add_run(
    "Dossier institutionnel pour l'introduction a UNDP :\n"
    "ordre de transmission et notes d'usage"
)
run.font.name = 'Georgia'
run.font.size = Pt(13)
run.font.color.rgb = NAVY

add_accent_rule(doc)

# ── Introduction ──
doc.add_paragraph(
    "Chere Professeure Scharff,\n\n"
    "Nous vous remercions sincerement de porter cette introduction aupres du "
    "Dr. Govinda Timilsina et de son equipe au UNDP. Le present document decrit "
    "les trois artefacts qui composent le dossier institutionnel SUNLIGHT, leur "
    "fonction respective dans la sequence d'engagement, et l'ordre dans lequel "
    "nous recommandons de les transmettre. Chaque document a ete concu pour servir "
    "un moment institutionnel precis ; leur transmission sequentielle permet de "
    "structurer l'engagement sans surcharger le destinataire initial."
)

doc.add_paragraph('')  # spacing

# ── Document 1 ──
p = doc.add_paragraph()
run = p.add_run("1. SUNLIGHT_Scharff_Note_Suivi_v2_2_FR.pdf")
run.bold = True
run.font.name = 'Georgia'
run.font.size = Pt(11)
run.font.color.rgb = NAVY

doc.add_paragraph(
    "Il s'agit de la note de suivi strategique qui prolonge la conversation initiee "
    "le 27 mars. Ce document presente l'architecture SUNLIGHT, son positionnement "
    "institutionnel vis-a-vis de l'ecosysteme d'integrite des marches publics du UNDP, "
    "et la proposition de collaboration structuree. C'est l'artefact qui ouvre "
    "l'engagement du Dr. Timilsina : il doit etre transmis avec l'email d'introduction "
    "comme piece jointe unique. Sa fonction est de susciter un interet institutionnel "
    "initial et de fournir le cadre strategique dans lequel les artefacts techniques "
    "ulterieurs prennent leur sens."
)

doc.add_paragraph('')

# ── Document 2 ──
p = doc.add_paragraph()
run = p.add_run("2. SUNLIGHT_Analytical_Dossier_v1.pdf")
run.bold = True
run.font.name = 'Georgia'
run.font.size = Pt(11)
run.font.color.rgb = NAVY

doc.add_paragraph(
    "Il s'agit du dossier analytique demontrant SUNLIGHT en operation : cinq contrats "
    "reels analyses a travers les trois moteurs du systeme, avec verdicts, scores "
    "dimensionnels, regles structurelles declenchees et citations juridiques rendues "
    "par le profil juridictionnel applicable. C'est l'artefact qui convertit un interet "
    "institutionnel initial en engagement technique substantiel. Il ne doit etre transmis "
    "que si et quand le Dr. Timilsina ou son equipe demande un examen plus approfondi "
    "ou une demonstration de la production analytique reelle du systeme."
)

doc.add_paragraph('')

# ── Document 3 ──
p = doc.add_paragraph()
run = p.add_run("3. SUNLIGHT_Institutional_FAQ_EN.pdf")
run.bold = True
run.font.name = 'Georgia'
run.font.size = Pt(11)
run.font.color.rgb = NAVY

doc.add_paragraph(
    "Il s'agit du document de reference couvrant les questions techniques et "
    "methodologiques previsibles que l'equipe du Dr. Timilsina est susceptible de "
    "soulever : integration avec Quantum, profil de faux positifs, modele tarifaire, "
    "propriete intellectuelle, deploiement en infrastructure UNDP, maintenance du "
    "standard MJPIS, et autres. Ce document peut etre transmis de maniere proactive "
    "aux cotes du dossier analytique si l'engagement technique est avance, ou de "
    "maniere reactive en reponse a des questions specifiques de la part du UNDP."
)

doc.add_paragraph('')
add_accent_rule(doc)

# ── Closing ──
doc.add_paragraph(
    "Nous mesurons pleinement le temps et le soin que vous consacrez a cette "
    "introduction, et nous vous en sommes profondement reconnaissants. Aucune "
    "pression de calendrier n'est placee de notre cote : la sequence de transmission "
    "est entierement a votre discretion, au rythme qui vous convient et qui convient "
    "a votre relation avec le Dr. Timilsina. Si vous avez besoin de materiel "
    "supplementaire, de clarifications, ou d'ajustements a l'un de ces documents "
    "avant transmission, nous restons a votre entiere disposition."
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(24)
p.add_run("Avec nos salutations respectueuses,").font.name = 'Georgia'

p = doc.add_paragraph()
run = p.add_run("Rimwaya Ouedraogo & Hugo Villalba")
run.font.name = 'Georgia'
run.italic = True

p = doc.add_paragraph()
run = p.add_run("Fondateurs, SUNLIGHT")
run.font.name = 'Georgia'
run.font.size = Pt(10)
run.font.color.rgb = DARK_GREY

# ── Save ──
doc.save(OUTPUT_PATH)
print(f"README saved to: {OUTPUT_PATH}")
